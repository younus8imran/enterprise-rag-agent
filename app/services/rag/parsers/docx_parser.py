"""DOCX file parser using python-docx."""

from pathlib import Path

from app.services.rag.parsers.base import BaseParser, ParsedDocument, ParsedPage


class DocxParser(BaseParser):
    """Parser for .docx files using python-docx."""

    async def parse(self, file_path: Path) -> ParsedDocument:
        """Parse a DOCX file, extracting paragraph text.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            ParsedDocument with the full document text as a single page.
        """
        self._validate_file(file_path)

        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install with: uv pip install python-docx"
            ) from exc

        doc = DocxDocument(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)

        return ParsedDocument(
            name=file_path.name,
            doc_type="docx",
            pages=[ParsedPage(text=full_text, page_number=1)],
            metadata={
                "source": str(file_path),
                "paragraph_count": len(paragraphs),
                "char_count": len(full_text),
            },
        )
